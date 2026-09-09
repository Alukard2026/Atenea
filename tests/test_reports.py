"""Informes y XLSX reales sobre PostgreSQL; todas las filas se revierten."""

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from sqlalchemy import update

from app.models import Client, Organization, Project, TimeEntry, User
from app.routers.reports import XLSX_TYPE
from app.worklog import monday_for
from test_worklog import WorklogCase


class ReportTests(WorklogCase):
    def report(self, month="2026-09"):
        response = self.browser.get("/reports/month", params={"month": month})
        self.assertEqual(response.status_code, 200)
        return response.context

    def excel(self, path):
        response = self.browser.get(path)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], XLSX_TYPE)
        self.assertEqual(response.headers["cache-control"], "no-store")
        book = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(book.close)
        return response, book

    def value_after_label(self, sheet, label):
        matches = [sheet.cell(cell.row, 6).value for row in sheet for cell in row if cell.value == label]
        self.assertEqual(len(matches), 1)
        return Decimal(str(matches[0]))

    def test_authenticated_default_month_and_navigation_links(self):
        response = self.browser.get("/reports/month")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["month"], date.today().strftime("%Y-%m"))
        for path in ("/dashboard", "/hours/history", "/hours/week"):
            self.assertIn('href="/reports/month"', self.browser.get(path).text)
        self.assertIn("/reports/week/export?week=", self.browser.get("/hours/week").text)

    def test_month_navigation_across_year(self):
        report = self.report("2026-01")
        self.assertEqual(report["previous_month"], "2025-12")
        self.assertEqual(report["next_month"], "2026-02")
        for month in (report["previous_month"], report["next_month"]):
            self.assertEqual(self.report(month)["month"], month)

    def test_unauthenticated_requests_redirect_to_login(self):
        self.browser.cookies.clear()
        for path in ("/reports/month", "/reports/month/export", "/reports/week/export"):
            response = self.browser.get(path)
            self.assertEqual(response.status_code, 303)
            self.assertEqual(response.headers["location"], "/login")

    def test_identity_parameters_and_duplicates_are_rejected(self):
        for path, field, value in (("/reports/month", "month", "2026-09"),
                                  ("/reports/month/export", "month", "2026-09"),
                                  ("/reports/week/export", "week", "2026-09-07")):
            for extra in ("user_id=1", "organization_id=1", "unexpected=1", f"{field}={value}"):
                self.assertEqual(self.browser.get(f"{path}?{field}={value}&{extra}").status_code, 400)

    def test_invalid_periods_are_safe_errors(self):
        for value in ("invalid", "2026-13", "1800-01", "2101-01", "2026-09-01"):
            for path in ("/reports/month", "/reports/month/export"):
                self.assertEqual(self.browser.get(path, params={"month": value}).status_code, 400)
        for value in ("invalid", "2026-02-30", "1800-01-01", "2101-01-01"):
            self.assertEqual(self.browser.get("/reports/week/export", params={"week": value}).status_code, 400)

    def test_month_boundaries_and_all_intersecting_weeks(self):
        self.add_entry(day=date(2026, 8, 31), hours="1.25")
        self.add_entry(day=date(2026, 9, 1), hours="2.50")
        self.add_entry(day=date(2026, 9, 6), hours="0.10")
        self.add_entry(day=date(2026, 9, 30), hours="3.75")
        self.add_entry(day=date(2026, 10, 1), hours="4.00")
        report = self.report()
        self.assertEqual(len(report["weeks"]), 5)
        self.assertEqual(report["weeks"][0], {"start": date(2026, 8, 31), "end": date(2026, 9, 6),
                                            "total": Decimal("3.85"), "month_total": Decimal("2.60")})
        self.assertEqual(report["weeks"][-1]["end"], date(2026, 10, 4))
        self.assertEqual(report["weeks"][1]["total"], Decimal("0.00"))
        self.assertEqual(report["active_days"], 3)
        self.assertEqual(report["total"], Decimal("6.35"))
        self.assertEqual(sum((week["month_total"] for week in report["weeks"]), Decimal("0.00")), report["total"])
        self.assertEqual(self.report("2026-08")["total"], Decimal("1.25"))
        self.assertEqual(self.report("2026-10")["total"], Decimal("4.00"))

    def test_six_week_month_and_leap_february(self):
        self.assertEqual(len(self.report("2026-08")["weeks"]), 6)
        self.add_entry(day=date(2024, 2, 29), hours="1.00")
        self.assertEqual(self.report("2024-02")["total"], Decimal("1.00"))
        self.assertEqual(self.report("2024-02")["end"], date(2024, 2, 29))

    def test_user_and_organization_isolation_in_html_and_both_exports(self):
        self.add_entry(day=date(2026, 9, 1), activity="Registro propio visible")
        self.add_entry(day=date(2026, 9, 1), user_id=self.colleague_id, hours="8", activity="CONFIDENCIAL COLEGA")
        self.add_entry(day=date(2026, 9, 1), user_id=self.foreign_user_id, organization_id=self.foreign_org_id,
                       client_id=self.foreign_client_id, project_id=self.foreign_project_id, hours="9", activity="CONFIDENCIAL EXTERNO")
        self.assertEqual(self.report()["total"], Decimal("2.25"))
        for path in ("/reports/month/export?month=2026-09", "/reports/week/export?week=2026-09-01"):
            _, book = self.excel(path)
            values = [cell.value for sheet in book for row in sheet for cell in row]
            self.assertIn("Registro propio visible", values)
            for private in ("CONFIDENCIAL COLEGA", "CONFIDENCIAL EXTERNO", "Cliente externo reservado", "Usuario externo"):
                self.assertNotIn(private, values)
        self.sign_in(self.colleague_id, self.org_id)
        self.assertEqual(self.report()["total"], Decimal("8.00"))
        _, book = self.excel("/reports/month/export?month=2026-09")
        self.assertNotIn("Registro propio visible", [cell.value for row in book["Detalle"] for cell in row])

    def test_active_status_is_revalidated_on_every_download(self):
        for model, identity in ((User, self.user_id), (Organization, self.org_id)):
            self.connection.execute(update(model).where(model.id == identity).values(is_active=False))
            for path in ("/reports/month", "/reports/week/export", "/reports/month/export"):
                self.sign_in(self.user_id, self.org_id)
                self.assertEqual(self.browser.get(path).headers.get("location"), "/login")
            self.connection.execute(update(model).where(model.id == identity).values(is_active=True))

    def test_empty_week_valid_workbook_and_seven_daily_totals(self):
        _, book = self.excel("/reports/week/export?week=2026-09-09")
        self.assertEqual(book.sheetnames, ["Semanal"])
        sheet = book["Semanal"]
        self.assertEqual(self.value_after_label(sheet, "Total semanal"), Decimal("0.00"))
        daily_rows = [row for row in sheet if isinstance(row[0].value, datetime)]
        self.assertEqual([row[0].value.date() for row in daily_rows], [date(2026, 9, 7) + timedelta(days=i) for i in range(7)])
        self.assertTrue(all(row[5].value == 0 and row[5].number_format == "0.00" for row in daily_rows))
        self.assertIsNone(sheet.auto_filter.ref)

    def test_empty_month_valid_workbook_and_zero_summary(self):
        self.assertEqual(self.report()["active_days"], 0)
        _, book = self.excel("/reports/month/export?month=2026-09")
        self.assertEqual(book.sheetnames, ["Resumen", "Detalle"])
        for name in book.sheetnames:
            for label in ("Total del mes", "Total facturable", "Total no facturable"):
                self.assertEqual(self.value_after_label(book[name], label), Decimal("0.00"))

    def test_week_filename_dates_and_cross_month_entries(self):
        self.add_entry(day=date(2026, 8, 30), activity="Fuera antes")
        self.add_entry(day=date(2026, 8, 31), activity="Lunes anterior")
        self.add_entry(day=date(2026, 9, 6), activity="Domingo incluido")
        self.add_entry(day=date(2026, 9, 7), activity="Fuera después")
        response, book = self.excel("/reports/week/export?week=2026-09-03")
        self.assertEqual(response.headers["content-disposition"], 'attachment; filename="Atenea_Horas_2026-08-31_al_2026-09-06.xlsx"')
        sheet = book.active
        self.assertEqual(sheet["B6"].value.date(), date(2026, 8, 31))
        self.assertEqual(sheet["D6"].value.date(), date(2026, 9, 6))
        self.assertIsInstance(sheet["F6"].value, datetime)
        values = [cell.value for row in sheet for cell in row]
        for included in ("Lunes anterior", "Domingo incluido"):
            self.assertIn(included, values)
        for excluded in ("Fuera antes", "Fuera después"):
            self.assertNotIn(excluded, values)
        self.assertEqual(self.value_after_label(sheet, "Total semanal"), Decimal("4.50"))

    def test_month_filename_detail_and_per_week_totals_exclude_adjacent_months(self):
        self.add_entry(day=date(2026, 8, 31), hours="7.00", activity="Mes anterior")
        self.add_entry(day=date(2026, 9, 1), hours="0.10")
        self.add_entry(day=date(2026, 9, 30), hours="0.20")
        self.add_entry(day=date(2026, 10, 1), hours="8.00", activity="Mes siguiente")
        response, book = self.excel("/reports/month/export?month=2026-09")
        self.assertEqual(response.headers["content-disposition"], 'attachment; filename="Atenea_Horas_2026-09.xlsx"')
        self.assertEqual(self.value_after_label(book["Resumen"], "Total por semana"), Decimal("0.30"))
        self.assertEqual(self.value_after_label(book["Resumen"], "Total por cliente"), Decimal("0.30"))
        weeks = [row for row in book["Resumen"] if isinstance(row[0].value, datetime)]
        self.assertEqual(sum(Decimal(str(row[5].value)) for row in weeks), Decimal("0.30"))
        for row in book["Detalle"]:
            if isinstance(row[0].value, datetime):
                self.assertEqual(row[0].value.month, 9)

    def test_decimal_billable_totals_and_days_count(self):
        self.add_entry(day=date(2026, 9, 1), hours="0.10")
        entry_id = self.add_entry(day=date(2026, 9, 1), hours="0.20")
        self.connection.execute(update(TimeEntry).where(TimeEntry.id == entry_id).values(billable=False))
        report = self.report()
        self.assertIsInstance(report["total"], Decimal)
        self.assertEqual(report["total"], Decimal("0.30"))
        self.assertEqual(report["billable_total"], Decimal("0.10"))
        self.assertEqual(report["non_billable_total"], Decimal("0.20"))
        self.assertEqual(report["active_days"], 1)
        for path, sheet_name in (("/reports/week/export?week=2026-09-01", "Semanal"), ("/reports/month/export?month=2026-09", "Resumen")):
            _, book = self.excel(path)
            self.assertEqual(self.value_after_label(book[sheet_name], "Total facturable"), Decimal("0.10"))
            self.assertEqual(self.value_after_label(book[sheet_name], "Total no facturable"), Decimal("0.20"))

    def test_workday_configuration_changes_classification_not_total(self):
        self.add_entry(day=date(2026, 9, 6), hours="1.50")
        _, book = self.excel("/reports/week/export?week=2026-09-01")
        self.assertEqual(self.value_after_label(book.active, "Total en días no laborables"), Decimal("1.50"))
        self.connection.execute(update(Organization).where(Organization.id == self.org_id).values(workday_sunday=True))
        _, book = self.excel("/reports/week/export?week=2026-09-01")
        self.assertEqual(self.value_after_label(book.active, "Total en días laborables"), Decimal("1.50"))
        self.assertEqual(self.value_after_label(book.active, "Total en días no laborables"), Decimal("0.00"))
        self.assertEqual(self.value_after_label(book.active, "Total semanal"), Decimal("1.50"))

    def test_optional_project_unicode_and_long_description(self):
        description = "Análisis jurídico: Ñandú, José, 日本語.\n" * 100
        entry_id = self.add_entry(day=date(2026, 9, 1))
        self.connection.execute(update(TimeEntry).where(TimeEntry.id == entry_id).values(project_id=None, description=description))
        for path, name in (("/reports/week/export?week=2026-09-01", "Semanal"), ("/reports/month/export?month=2026-09", "Detalle")):
            _, book = self.excel(path)
            self.assertEqual(book[name]["D9"].value, "Sin proyecto")
            self.assertEqual(book[name]["E9"].value, description)
            self.assertTrue(book[name]["E9"].alignment.wrap_text)

    def test_formula_injection_all_user_text_is_literal(self):
        entry_id = self.add_entry(day=date(2026, 9, 1))
        for prefix in ("=", "+", "-", "@", " \t="):
            malicious = prefix + 'HYPERLINK("https://example.invalid","texto")'
            for model, identity, field in ((Organization, self.org_id, "name"), (User, self.user_id, "full_name"),
                                           (Client, self.client_id, "name"), (Project, self.project_id, "name"),
                                           (TimeEntry, entry_id, "description")):
                self.connection.execute(update(model).where(model.id == identity).values(**{field: malicious}))
            for path in ("/reports/week/export?week=2026-09-01", "/reports/month/export?month=2026-09"):
                response, book = self.excel(path)
                found = []
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, "f")
                            self.assertIsNone(cell.hyperlink)
                            if cell.value == malicious:
                                found.append(cell)
                                self.assertEqual(cell.data_type, "s")
                                self.assertTrue(cell.quotePrefix)
                self.assertGreaterEqual(len(found), 5)
                with ZipFile(BytesIO(response.content)) as archive:
                    self.assertFalse(any("vba" in name.lower() or "externallinks" in name.lower() for name in archive.namelist()))

    def test_invalid_xml_controls_do_not_break_workbooks(self):
        entry_id = self.add_entry(day=date(2026, 9, 1))
        self.connection.execute(update(TimeEntry).where(TimeEntry.id == entry_id).values(description="Texto\x01válido\uffff"))
        _, book = self.excel("/reports/week/export?week=2026-09-01")
        self.assertEqual(book.active["E9"].value, "Texto\ufffdválido\ufffd")

    def test_historical_archived_clients_are_exported(self):
        self.add_entry(day=date(2026, 9, 1))
        self.connection.execute(update(Client).where(Client.id == self.client_id).values(is_active=False))
        self.connection.execute(update(Project).where(Project.id == self.project_id).values(is_active=False))
        self.assertEqual(self.report()["total"], Decimal("2.25"))
        _, book = self.excel("/reports/month/export?month=2026-09")
        self.assertEqual(book["Detalle"]["C9"].value, "Cliente principal")

    def test_excel_format_dates_numbers_filter_and_print(self):
        self.add_entry(day=date(2026, 9, 1), hours="2.25")
        for path in ("/reports/week/export?week=2026-09-01", "/reports/month/export?month=2026-09"):
            _, book = self.excel(path)
            for sheet in book:
                self.assertTrue(sheet.freeze_panes)
                self.assertFalse(sheet.sheet_view.showGridLines)
                self.assertEqual(sheet.page_setup.orientation, "landscape")
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                self.assertEqual(sheet.page_setup.fitToHeight, 0)
                self.assertTrue(sheet.print_area)
                self.assertGreater(sheet.page_margins.left, 0)
            detail = book["Detalle"] if "Detalle" in book.sheetnames else book.active
            self.assertEqual(detail.auto_filter.ref, "A8:G9")
            self.assertEqual(detail["A9"].number_format, "dd/mm/yyyy")
            self.assertEqual(detail["F9"].number_format, "0.00")
            self.assertEqual(detail["F9"].data_type, "n")
            self.assertTrue(detail["A8"].font.bold)
            self.assertEqual(detail["G9"].value, "Sí")

    def test_export_defaults_match_current_period(self):
        response, book = self.excel("/reports/week/export")
        self.assertEqual(book.active["B6"].value.date(), monday_for(date.today()))
        response, book = self.excel("/reports/month/export")
        self.assertIn(date.today().strftime("%Y-%m"), response.headers["content-disposition"])
        self.assertEqual(book["Resumen"]["B6"].value.date(), date.today().replace(day=1))

    def test_supported_calendar_limits(self):
        self.assertIsNone(self.report("1900-01")["previous_month"])
        self.assertIsNone(self.report("2100-12")["next_month"])
        for month in ("1900-01", "2100-12"):
            self.excel("/reports/month/export?month=" + month)
